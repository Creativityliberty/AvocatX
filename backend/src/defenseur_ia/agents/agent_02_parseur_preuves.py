"""
AGENT 2 - Parseur de Preuves
Analyse et extrait le contenu des pièces justificatives (PDF, images, emails, etc.)
"""

import asyncio
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
import json

# Importations optionnelles (graceful fallbacks si non disponibles)
try:
    import pytesseract  # type: ignore
except Exception:  # pragma: no cover - environnement sans OCR
    pytesseract = None  # type: ignore

try:
    from PIL import Image  # type: ignore
except Exception:  # pragma: no cover
    Image = None  # type: ignore

try:
    import fitz  # type: ignore  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None  # type: ignore

try:
    import mailparser  # type: ignore
except Exception:  # pragma: no cover
    mailparser = None  # type: ignore

try:
    from docx import Document  # type: ignore
except Exception:  # pragma: no cover
    Document = None  # type: ignore

from ..core.base import BaseNode, NodeContext
from ..models import PieceParsed, MetaInfo
from ..services.storage_service import StorageService

logger = logging.getLogger(__name__)

class ParseurPreuvesNode(BaseNode):
    """
    Agent responsable du parsing des pièces justificatives
    
    Inputs: pieces (liste des fichiers uploadés)
    Output: pieces_parsed (contenu extrait et structuré)
    """
    
    def __init__(self):
        super().__init__("parseur_preuves", "Parseur de Preuves")
        self.storage_service = StorageService()
        
    async def exec(self, context: NodeContext) -> Dict[str, Any]:
        """
        Parse toutes les pièces justificatives uploadées
        """
        try:
            logger.info("🔍 Début du parsing des pièces justificatives")
            
            # Récupération des métadonnées et pièces
            meta_info = context.shared_store.get("meta_info", {})
            pieces_files = context.shared_store.get("pieces", [])
            
            if not pieces_files:
                logger.warning("Aucune pièce à parser")
                return {"pieces_parsed": []}
            
            pieces_parsed = []
            
            for piece_info in pieces_files:
                try:
                    parsed_piece = await self._parse_single_piece(piece_info)
                    if parsed_piece:
                        pieces_parsed.append(parsed_piece)
                        logger.info(f"✅ Pièce parsée: {piece_info.get('filename', 'unknown')}")
                        
                except Exception as e:
                    logger.error(f"❌ Erreur parsing {piece_info.get('filename', 'unknown')}: {e}")
                    # Créer une entrée d'erreur
                    error_piece = PieceParsed(
                        filename=piece_info.get('filename', 'unknown'),
                        type_piece="error",
                        contenu_brut="Erreur lors du parsing",
                        contenu_structure={},
                        metadata={
                            "error": str(e),
                            "original_path": piece_info.get('path', '')
                        }
                    )
                    pieces_parsed.append(error_piece.dict())
            
            # Sauvegarde dans le store
            context.shared_store.set("pieces_parsed", pieces_parsed)
            
            # Statistiques
            total_pieces = len(pieces_files)
            parsed_success = len([p for p in pieces_parsed if p.get('type_piece') != 'error'])
            
            logger.info(f"📊 Parsing terminé: {parsed_success}/{total_pieces} pièces parsées avec succès")
            
            return {
                "pieces_parsed": pieces_parsed,
                "stats": {
                    "total": total_pieces,
                    "success": parsed_success,
                    "errors": total_pieces - parsed_success
                }
            }
            
        except Exception as e:
            logger.error(f"❌ Erreur critique dans le parseur de preuves: {e}")
            raise
    
    async def _parse_single_piece(self, piece_info: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        """Parse une seule pièce justificative"""
        filename = piece_info.get('filename', '')
        file_path = piece_info.get('path', '')
        
        if not file_path or not Path(file_path).exists():
            logger.error(f"Fichier introuvable: {file_path}")
            return None
        
        # Détection du type de fichier
        file_extension = Path(filename).suffix.lower()
        
        if file_extension == '.pdf':
            return await self._parse_pdf(file_path, filename)
        elif file_extension in ['.jpg', '.jpeg', '.png', '.tiff', '.bmp']:
            return await self._parse_image(file_path, filename)
        elif file_extension == '.eml':
            return await self._parse_email(file_path, filename)
        elif file_extension == '.docx':
            return await self._parse_docx(file_path, filename)
        elif file_extension == '.txt':
            return await self._parse_text(file_path, filename)
        else:
            logger.warning(f"Type de fichier non supporté: {file_extension}")
            return self._create_unsupported_piece(filename, file_path)
    
    async def _parse_pdf(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse un fichier PDF (PyMuPDF si dispo, sinon fallback pdfplumber)"""
        try:
            contenu_brut = ""
            pages_info: List[Dict[str, Any]] = []

            if fitz is not None:
                # Chemin rapide via PyMuPDF
                doc = fitz.open(file_path)
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    text = page.get_text()
                    contenu_brut += f"\n--- Page {page_num + 1} ---\n{text}"
                    pages_info.append({
                        "page": page_num + 1,
                        "text_length": len(text),
                        "has_images": len(page.get_images()) > 0,
                    })
                total_pages = len(doc)
                parsing_method = "PyMuPDF"
                doc.close()
            else:
                # Fallback: pdfplumber
                try:
                    import pdfplumber  # type: ignore
                except Exception as e:
                    logger.warning(f"pdfplumber indisponible: {e}")
                    pdfplumber = None  # type: ignore

                if pdfplumber is None:
                    logger.warning("Aucune librairie PDF disponible (PyMuPDF/pdfplumber)")
                    contenu_brut = ""
                    total_pages = 0
                    parsing_method = "none"
                else:
                    with pdfplumber.open(file_path) as pdf:
                        for i, page in enumerate(pdf.pages):
                            text = page.extract_text() or ""
                            contenu_brut += f"\n--- Page {i + 1} ---\n{text}"
                            pages_info.append({
                                "page": i + 1,
                                "text_length": len(text),
                                "has_images": False,
                            })
                        total_pages = len(pdf.pages)
                    parsing_method = "pdfplumber"

            type_piece = self._detect_document_type(contenu_brut, filename)
            contenu_structure = {
                "pages": pages_info,
                "total_pages": total_pages,
                "text_length": len(contenu_brut),
                "detected_type": type_piece,
            }
            entities = self._extract_entities(contenu_brut)
            contenu_structure.update(entities)
            return PieceParsed(
                filename=filename,
                type_piece=type_piece,
                contenu_brut=contenu_brut,
                contenu_structure=contenu_structure,
                metadata={
                    "file_type": "pdf",
                    "original_path": file_path,
                    "parsing_method": parsing_method,
                },
            ).dict()

        except Exception as e:
            logger.error(f"Erreur parsing PDF {filename}: {e}")
            raise
    
    async def _parse_image(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse une image avec OCR (si disponible)"""
        try:
            parsing_method = "none"
            contenu_brut = ""
            width = height = 0
            mode = ""

            if Image is not None:
                image = Image.open(file_path)
                width, height = image.size
                mode = image.mode
                if pytesseract is not None:
                    try:
                        contenu_brut = pytesseract.image_to_string(image, lang='fra')
                        parsing_method = "Tesseract OCR"
                    except Exception as ocr_err:
                        logger.warning(f"OCR indisponible ou erreur: {ocr_err}")
                        parsing_method = "image_only"
                else:
                    logger.warning("pytesseract non installé - OCR désactivé")
                    parsing_method = "image_only"
            else:
                logger.warning("Pillow (PIL) non installé - impossible de lire l'image")

            type_piece = self._detect_document_type(contenu_brut, filename)
            contenu_structure = {
                "image_width": width,
                "image_height": height,
                "image_mode": mode,
                "text_length": len(contenu_brut),
                "detected_type": type_piece,
            }
            entities = self._extract_entities(contenu_brut)
            contenu_structure.update(entities)

            return PieceParsed(
                filename=filename,
                type_piece=type_piece,
                contenu_brut=contenu_brut,
                contenu_structure=contenu_structure,
                metadata={
                    "file_type": "image",
                    "original_path": file_path,
                    "parsing_method": parsing_method,
                },
            ).dict()

        except Exception as e:
            logger.error(f"Erreur parsing image {filename}: {e}")
            raise
    
    async def _parse_email(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse un email .eml (mailparser si dispo, sinon fallback lecture brute)"""
        try:
            parsing_method = "mailparser" if mailparser is not None else "raw_read"
            if mailparser is not None:
                mail = mailparser.parse_from_file(file_path)
                contenu_brut = f"""
De: {mail.from_}
À: {', '.join(mail.to) if mail.to else 'N/A'}
Sujet: {mail.subject or 'N/A'}
Date: {mail.date or 'N/A'}

{mail.body or 'Pas de contenu'}
"""
                contenu_structure = {
                    "from": mail.from_,
                    "to": mail.to,
                    "subject": mail.subject,
                    "date": str(mail.date) if mail.date else None,
                    "attachments": len(mail.attachments) if mail.attachments else 0,
                    "detected_type": "email",
                }
            else:
                try:
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        contenu_brut = f.read()
                except Exception:
                    contenu_brut = ""
                contenu_structure = {
                    "from": None,
                    "to": None,
                    "subject": None,
                    "date": None,
                    "attachments": 0,
                    "detected_type": "email",
                }

            entities = self._extract_entities(contenu_brut)
            contenu_structure.update(entities)
            return PieceParsed(
                filename=filename,
                type_piece="email",
                contenu_brut=contenu_brut,
                contenu_structure=contenu_structure,
                metadata={
                    "file_type": "email",
                    "original_path": file_path,
                    "parsing_method": parsing_method,
                },
            ).dict()

        except Exception as e:
            logger.error(f"Erreur parsing email {filename}: {e}")
            raise
    
    async def _parse_docx(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse un document Word (python-docx si dispo)"""
        try:
            parsing_method = "python-docx" if Document is not None else "raw_read"
            contenu_brut = ""
            paragraphs_count = 0

            if Document is not None:
                doc = Document(file_path)
                for paragraph in doc.paragraphs:
                    contenu_brut += paragraph.text + "\n"
                paragraphs_count = len(doc.paragraphs)
            else:
                try:
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        contenu_brut = f.read()
                except Exception:
                    contenu_brut = ""

            type_piece = self._detect_document_type(contenu_brut, filename)
            contenu_structure = {
                "paragraphs": paragraphs_count,
                "text_length": len(contenu_brut),
                "detected_type": type_piece,
            }
            entities = self._extract_entities(contenu_brut)
            contenu_structure.update(entities)
            return PieceParsed(
                filename=filename,
                type_piece=type_piece,
                contenu_brut=contenu_brut,
                contenu_structure=contenu_structure,
                metadata={
                    "file_type": "docx",
                    "original_path": file_path,
                    "parsing_method": parsing_method,
                },
            ).dict()

        except Exception as e:
            logger.error(f"Erreur parsing DOCX {filename}: {e}")
            raise
    
    async def _parse_text(self, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse un fichier texte simple"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                contenu_brut = f.read()
            
            type_piece = self._detect_document_type(contenu_brut, filename)
            
            contenu_structure = {
                "text_length": len(contenu_brut),
                "lines": len(contenu_brut.split('\n')),
                "detected_type": type_piece
            }
            
            # Extraction d'entités
            entities = self._extract_entities(contenu_brut)
            contenu_structure.update(entities)
            
            return PieceParsed(
                filename=filename,
                type_piece=type_piece,
                contenu_brut=contenu_brut,
                contenu_structure=contenu_structure,
                metadata={
                    "file_type": "text",
                    "original_path": file_path,
                    "parsing_method": "direct_read"
                }
            ).dict()
            
        except Exception as e:
            logger.error(f"Erreur parsing texte {filename}: {e}")
            raise
    
    def _detect_document_type(self, content: str, filename: str) -> str:
        """Détecte le type de document basé sur le contenu"""
        content_lower = content.lower()
        filename_lower = filename.lower()
        
        # Types de documents administratifs
        if any(keyword in content_lower for keyword in ['oqtf', 'obligation de quitter', 'préfecture']):
            return "decision_administrative"
        elif any(keyword in content_lower for keyword in ['passeport', 'carte d\'identité', 'permis']):
            return "piece_identite"
        elif any(keyword in content_lower for keyword in ['contrat de travail', 'salaire', 'employeur']):
            return "document_travail"
        elif any(keyword in content_lower for keyword in ['bail', 'loyer', 'logement']):
            return "justificatif_domicile"
        elif any(keyword in content_lower for keyword in ['médical', 'certificat médical', 'hôpital']):
            return "document_medical"
        elif any(keyword in content_lower for keyword in ['mariage', 'naissance', 'état civil']):
            return "acte_etat_civil"
        elif 'email' in filename_lower or '@' in content:
            return "email"
        else:
            return "document_autre"
    
    def _extract_entities(self, content: str) -> Dict[str, Any]:
        """Extraction d'entités importantes du contenu"""
        import re
        
        entities = {
            "dates": [],
            "noms": [],
            "adresses": [],
            "numeros": []
        }
        
        # Extraction de dates (format français)
        date_patterns = [
            r'\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{4}',
            r'\d{1,2}\s+\w+\s+\d{4}'
        ]
        
        for pattern in date_patterns:
            dates = re.findall(pattern, content)
            entities["dates"].extend(dates)
        
        # Extraction de numéros (téléphone, SIRET, etc.)
        numero_patterns = [
            r'\b\d{10}\b',  # Téléphone
            r'\b\d{14}\b',  # SIRET
            r'\b\d{9}\b'    # SIREN
        ]
        
        for pattern in numero_patterns:
            numeros = re.findall(pattern, content)
            entities["numeros"].extend(numeros)
        
        # Nettoyage des doublons
        for key in entities:
            entities[key] = list(set(entities[key]))
        
        return entities
    
    def _create_unsupported_piece(self, filename: str, file_path: str) -> Dict[str, Any]:
        """Crée une entrée pour un type de fichier non supporté"""
        return PieceParsed(
            filename=filename,
            type_piece="non_supporte",
            contenu_brut="Type de fichier non supporté pour le parsing automatique",
            contenu_structure={"error": "unsupported_file_type"},
            metadata={
                "file_type": "unsupported",
                "original_path": file_path,
                "parsing_method": "none"
            }
        ).dict()
